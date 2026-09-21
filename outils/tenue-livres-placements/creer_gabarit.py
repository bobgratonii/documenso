#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Crée le gabarit de saisie Excel (un fichier par client, un onglet par compte).

    python creer_gabarit.py                       -> Gabarit_Client_vide.xlsx (vide, prêt à remplir)
    python creer_gabarit.py --demo                -> Gabarit_Client_Exemple.xlsx + taux_fictifs_demo.csv
                                                     avec des données FICTIVES pour tester

Toutes les données de la démo sont inventées : titres, prix, quantités et taux de change.
"""
from __future__ import annotations

import argparse
import csv
import math
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from calcul_placements import (  # noqa: E402
    COLONNES_1099, COLONNES_COMPTES, COLONNES_RELEVES, COLONNES_TX, ENCAISSE, SERIE_ANNUELLE,
    SERIE_QUOTIDIENNE, TYPES_VALIDES, Compte, Ligne1099, LigneReleve, Parametres, Taux, Transaction,
    calculer, charger_taux_csv, lire_classeur,
)

ENTETE_FILL = PatternFill("solid", fgColor="1F3864")
ENTETE_FONT = Font(bold=True, color="FFFFFF")
SAISIE_FILL = PatternFill("solid", fgColor="FFF2CC")
NB_LIGNES_SAISIE = 500

INSTRUCTIONS = [
    ("GABARIT DE TENUE DE LIVRES : PLACEMENTS AMÉRICAINS D'UN CONTRIBUABLE CANADIEN", True),
    ("", False),
    ("Un fichier par client (même contribuable). Un onglet par compte, nommé exactement comme le code dans l'onglet Comptes.", False),
    ("Les cellules jaunes se remplissent à la main à partir des relevés papier. Le script calcul_placements.py fait le reste.", False),
    ("Le PBR est regroupé par titre pour le contribuable (tous comptes confondus), conformément à l'article 47 LIR.", False),
    ("", False),
    ("TYPES D'OPÉRATION", True),
    ("ACHAT : Quantité, Prix USD, Commission USD. Le montant brut se calcule (quantité × prix) si la colonne est vide.", False),
    ("VENTE : Quantité, Prix USD, Commission USD. Produit net = brut − commission.", False),
    ("DIVIDENDE : Titre, Montant brut USD, Retenue USD (impôt américain retenu, normalement 15 % avec W-8BEN).", False),
    ("INTERET : Titre = ENCAISSE (ou le titre obligataire), Montant brut USD, Retenue USD si applicable.", False),
    ("REMB_CAPITAL : remboursement de capital ; réduit le PBR. Montant brut USD.", False),
    ("FRAIS : honoraires de gestion, frais de garde. Titre = ENCAISSE, Montant brut USD.", False),
    ("FRACTIONNEMENT : Quantité = titres AJOUTÉS (2 pour 1 sur 100 titres → 100 ; regroupement → quantité négative).", False),
    ("TRANSFERT_TITRES : saisir UNE seule fois, dans l'onglet du compte SOURCE, avec Quantité et Compte contrepartie. PBR inchangé.", False),
    ("TRANSFERT_ENCAISSE : saisir UNE seule fois, dans le compte SOURCE, Montant brut USD et Compte contrepartie.", False),
    ("DEPOT / RETRAIT : entrée ou sortie d'argent du client. Titre = ENCAISSE, Montant brut USD.", False),
    ("AJUST_PBR : ajustement manuel du PBR (distribution fantôme, correction). Montant brut USD, signe + ou −.", False),
    ("SOLDE_OUVERTURE : positions au début du suivi. Titre, Quantité, Montant brut USD = coût USD, PBR CAD (ouverture) = PBR CAD connu.", False),
    ("    Pour l'encaisse d'ouverture : Titre = ENCAISSE, Montant brut USD = solde.", False),
    ("Dividende réinvesti : saisir deux lignes, un DIVIDENDE puis un ACHAT du même montant.", False),
    ("", False),
    ("DATES ET TAUX", True),
    ("Date transaction obligatoire. Date règlement facultative. L'onglet Paramètres choisit laquelle sert au taux de change.", False),
    ("Taux quotidien de la Banque du Canada à la date de l'opération. Fin de semaine ou férié : jour ouvrable précédent.", False),
    ("Revenus : taux quotidien ou moyenne annuelle BdC (choix dans Paramètres). Dispositions : toujours le taux quotidien.", False),
    ("", False),
    ("CONTRÔLES", True),
    ("Relevés : à chaque fin de mois, copier la quantité par titre et le solde d'encaisse USD du relevé. Le script compare.", False),
    ("1099 : à la fin de l'année, copier les cases du 1099 composite. Le côté USD du calcul doit tomber juste.", False),
    ("Attention : le coût du 1099-B est en USD selon les règles américaines et ne vérifie jamais le PBR canadien.", False),
    ("", False),
    ("HORS DU GABARIT", True),
    ("Placements privés ou sociétés en commandite avec K-1, fonds monétaires, gains de change sur l'encaisse (39(2)) : à traiter à part.", False),
]


def _entete(ws, colonnes, largeurs=None):
    for j, h in enumerate(colonnes, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.fill = ENTETE_FILL
        c.font = ENTETE_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = (largeurs or {}).get(j, max(12, len(h) + 3))
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def _zone_saisie(ws, nb_colonnes, formats=None):
    for i in range(2, NB_LIGNES_SAISIE + 2):
        for j in range(1, nb_colonnes + 1):
            c = ws.cell(row=i, column=j)
            c.fill = SAISIE_FILL
            if formats and j in formats:
                c.number_format = formats[j]


def _onglet_compte(wb, code):
    ws = wb.create_sheet(code)
    _entete(ws, COLONNES_TX, {3: 20, 4: 12, 5: 34, 11: 20, 12: 20, 13: 40})
    _zone_saisie(ws, len(COLONNES_TX), {1: "yyyy-mm-dd", 2: "yyyy-mm-dd", 6: "#,##0.####", 7: "#,##0.00",
                                        8: "#,##0.00", 9: "#,##0.00", 10: "#,##0.00", 12: "#,##0.00"})
    dv = DataValidation(type="list", formula1='"' + ",".join(TYPES_VALIDES) + '"', allow_blank=True,
                        errorTitle="Type", error="Choisir un type dans la liste")
    ws.add_data_validation(dv)
    dv.add(f"C2:C{NB_LIGNES_SAISIE + 1}")
    dv2 = DataValidation(type="list", formula1=f"=Comptes!$A$2:$A$50", allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add(f"K2:K{NB_LIGNES_SAISIE + 1}")
    dv3 = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                         error="Entrer une date (AAAA-MM-JJ)")
    ws.add_data_validation(dv3)
    dv3.add(f"A2:B{NB_LIGNES_SAISIE + 1}")
    return ws


def creer_classeur(chemin: Path, client: str, annee: int, comptes: list[Compte], convention="TRANSACTION",
                   methode="QUOTIDIEN") -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    for i, (texte, gras) in enumerate(INSTRUCTIONS, start=1):
        c = ws.cell(row=i, column=1, value=texte)
        if gras:
            c.font = Font(bold=True, size=12 if i == 1 else 11)
    ws.column_dimensions["A"].width = 150

    ws = wb.create_sheet("Paramètres")
    lignes = [
        ("Client", client, "Nom du contribuable (un seul contribuable par fichier)"),
        ("Année", annee, "Année d'imposition à produire ; l'historique antérieur sert au PBR"),
        ("Convention de date", convention, "TRANSACTION ou REGLEMENT : date qui détermine le taux de change"),
        ("Taux pour les revenus", methode, "QUOTIDIEN ou MOYENNE_ANNUELLE (dividendes et intérêts seulement)"),
    ]
    _entete(ws, ["Paramètre", "Valeur", "Explication"], {1: 24, 2: 24, 3: 70})
    for i, (k, v, e) in enumerate(lignes, start=2):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v).fill = SAISIE_FILL
        ws.cell(row=i, column=3, value=e)
    dvc = DataValidation(type="list", formula1='"TRANSACTION,REGLEMENT"')
    ws.add_data_validation(dvc)
    dvc.add("B4")
    dvm = DataValidation(type="list", formula1='"QUOTIDIEN,MOYENNE_ANNUELLE"')
    ws.add_data_validation(dvm)
    dvm.add("B5")

    ws = wb.create_sheet("Comptes")
    _entete(ws, COLONNES_COMPTES, {1: 14, 2: 30, 3: 20, 4: 40})
    _zone_saisie(ws, len(COLONNES_COMPTES))
    for i, c in enumerate(comptes, start=2):
        ws.cell(row=i, column=1, value=c.code)
        ws.cell(row=i, column=2, value=c.courtier)
        ws.cell(row=i, column=3, value=c.numero)
        ws.cell(row=i, column=4, value=c.description)

    for c in comptes:
        _onglet_compte(wb, c.code)

    ws = wb.create_sheet("Relevés")
    _entete(ws, COLONNES_RELEVES, {1: 14, 2: 14, 3: 14, 4: 22})
    _zone_saisie(ws, len(COLONNES_RELEVES), {2: "yyyy-mm-dd", 4: "#,##0.####"})
    dv = DataValidation(type="list", formula1="=Comptes!$A$2:$A$50", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"A2:A{NB_LIGNES_SAISIE + 1}")

    ws = wb.create_sheet("1099")
    _entete(ws, COLONNES_1099, {1: 14, 2: 10, 3: 22, 4: 26, 5: 16, 6: 26})
    _zone_saisie(ws, len(COLONNES_1099), {3: "#,##0.00", 4: "#,##0.00", 5: "#,##0.00", 6: "#,##0.00"})
    dv = DataValidation(type="list", formula1="=Comptes!$A$2:$A$50", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"A2:A{NB_LIGNES_SAISIE + 1}")

    wb.save(chemin)
    return wb


# ---------------------------------------------------------------------------
# Démo avec données fictives
# ---------------------------------------------------------------------------

def taux_fictifs(debut: date, fin: date) -> Taux:
    """Série FICTIVE et déterministe autour de 1.36, jours ouvrables seulement."""
    q = {}
    d = debut
    k = 0
    while d <= fin:
        if d.weekday() < 5:
            q[d] = round(1.36 + 0.05 * math.sin(k / 37) + 0.012 * math.cos(k / 5.3), 4)
            k += 1
        d += timedelta(days=1)
    t = Taux(q)
    for a in range(debut.year, fin.year + 1):
        vals = [v for dd, v in q.items() if dd.year == a]
        t.annuels[a] = round(sum(vals) / len(vals), 4)
    return t


def ecrire_taux_csv(chemin: Path, t: Taux) -> None:
    with open(chemin, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["date", "serie", "valeur"])
        for d in t.dates:
            w.writerow([d.isoformat(), SERIE_QUOTIDIENNE, f"{t.quotidiens[d]:.4f}"])
        for a in sorted(t.annuels):
            w.writerow([f"{a}-12-31", SERIE_ANNUELLE, f"{t.annuels[a]:.4f}"])


def transactions_demo() -> dict[str, list[list]]:
    """Lignes fictives par compte, dans l'ordre des colonnes COLONNES_TX."""
    D = date
    return {
        "RCM-1001": [
            [D(2024, 1, 2), None, "SOLDE_OUVERTURE", ENCAISSE, "Encaisse USD au début du suivi", None, None, 135000, None, None, None, None, "Solde selon relevé du 31 déc. 2023"],
            [D(2024, 1, 2), None, "SOLDE_OUVERTURE", "AAPL", "Apple Inc.", 200, None, 30000, None, None, None, 39600, "PBR CAD selon dossier 2023"],
            [D(2024, 3, 5), D(2024, 3, 7), "ACHAT", "MSFT", "Microsoft Corp.", 100, 405.00, None, 9.95, None, None, None, None],
            [D(2024, 3, 15), None, "DIVIDENDE", "AAPL", "Apple Inc.", None, None, 48.00, None, 7.20, None, None, None],
            [D(2024, 6, 10), D(2024, 6, 12), "ACHAT", "VTI", "Vanguard Total Stock Market ETF", 300, 260.00, None, 9.95, None, None, None, None],
            [D(2024, 6, 28), None, "DIVIDENDE", "VTI", "Vanguard Total Stock Market ETF", None, None, 270.00, None, 40.50, None, None, None],
            [D(2024, 9, 20), D(2024, 9, 24), "VENTE", "AAPL", "Apple Inc.", 50, 228.00, None, 9.95, None, None, None, None],
            [D(2024, 12, 31), None, "INTERET", ENCAISSE, "Intérêts sur encaisse", None, None, 310.25, None, None, None, None, None],
            [D(2025, 1, 10), None, "FRAIS", ENCAISSE, "Honoraires de gestion T4 2024", None, None, 1250.00, None, None, None, None, None],
            [D(2025, 2, 14), None, "DIVIDENDE", "AAPL", "Apple Inc.", None, None, 37.50, None, 5.63, None, None, None],
            [D(2025, 3, 3), D(2025, 3, 5), "VENTE", "MSFT", "Microsoft Corp.", 100, 380.00, None, 9.95, None, None, None, None],
            [D(2025, 3, 20), D(2025, 3, 24), "ACHAT", "MSFT", "Microsoft Corp.", 60, 385.00, None, 9.95, None, None, None, "Rachat dans les 30 jours : perte apparente partielle"],
            [D(2025, 4, 1), None, "TRANSFERT_TITRES", "VTI", "Vanguard Total Stock Market ETF", 100, None, None, None, None, "RCM-1002", None, "Transfert en nature vers le compte 1002"],
            [D(2025, 4, 1), None, "TRANSFERT_ENCAISSE", ENCAISSE, "Virement interne", None, None, 10000, None, None, "RCM-1002", None, None],
            [D(2025, 5, 15), None, "REMB_CAPITAL", "VTI", "Vanguard Total Stock Market ETF", None, None, 45.00, None, None, None, None, "Fictif : remboursement de capital"],
            [D(2025, 6, 27), None, "DIVIDENDE", "VTI", "Vanguard Total Stock Market ETF", None, None, 190.00, None, 28.50, None, None, None],
            [D(2025, 7, 15), None, "FRACTIONNEMENT", "AAPL", "Apple Inc.", 150, None, None, None, None, None, None, "Fictif : 2 pour 1 sur 150 titres"],
            [D(2025, 8, 12), D(2025, 8, 14), "VENTE", "AAPL", "Apple Inc.", 100, 115.00, None, 9.95, None, None, None, None],
            [D(2025, 9, 30), None, "DEPOT", ENCAISSE, "Virement du client", None, None, 50000, None, None, None, None, None],
            [D(2025, 11, 14), None, "DIVIDENDE", "AAPL", "Apple Inc.", None, None, 52.00, None, 15.60, None, None, "Fictif : retenue de 30 % pour montrer l'alerte"],
            [D(2025, 12, 31), None, "INTERET", ENCAISSE, "Intérêts sur encaisse", None, None, 412.80, None, None, None, None, None],
        ],
        "RCM-1002": [
            [D(2024, 1, 2), None, "SOLDE_OUVERTURE", ENCAISSE, "Encaisse USD au début du suivi", None, None, 30000, None, None, None, None, None],
            [D(2024, 5, 2), D(2024, 5, 6), "ACHAT", "VTI", "Vanguard Total Stock Market ETF", 100, 255.00, None, 9.95, None, None, None, None],
            [D(2024, 6, 28), None, "DIVIDENDE", "VTI", "Vanguard Total Stock Market ETF", None, None, 90.00, None, 13.50, None, None, None],
            [D(2025, 6, 27), None, "DIVIDENDE", "VTI", "Vanguard Total Stock Market ETF", None, None, 190.00, None, 28.50, None, None, None],
            [D(2025, 10, 6), D(2025, 10, 8), "VENTE", "VTI", "Vanguard Total Stock Market ETF", 150, 300.00, None, 9.95, None, None, None, "PBR regroupé avec le compte 1001"],
            [D(2025, 11, 3), None, "RETRAIT", ENCAISSE, "Virement au client", None, None, 20000, None, None, None, None, None],
        ],
    }


def demo(dossier: Path) -> tuple[Path, Path]:
    comptes = [
        Compte("RCM-1001", "Rockefeller (fictif)", "XXXX-1001", "Compte principal"),
        Compte("RCM-1002", "Rockefeller (fictif)", "XXXX-1002", "Compte secondaire"),
    ]
    chemin = dossier / "Gabarit_Client_Exemple.xlsx"
    chemin_taux = dossier / "taux_fictifs_demo.csv"
    t = taux_fictifs(date(2023, 11, 1), date(2026, 12, 31))
    ecrire_taux_csv(chemin_taux, t)

    wb = creer_classeur(chemin, "Client Exemple (données fictives)", 2025, comptes)
    for code, lignes in transactions_demo().items():
        ws = wb[code]
        for i, l in enumerate(lignes, start=2):
            for j, v in enumerate(l, start=1):
                ws.cell(row=i, column=j, value=v)
    wb.save(chemin)

    # Relevés mensuels et 1099 fictifs, rendus cohérents avec les opérations en les calculant
    params, comptes_lus, txs, _, _, _ = lire_classeur(chemin)
    res = calculer(txs, comptes_lus, t, params)
    fins_de_mois = []
    for a in (2024, 2025):
        for m in range(1, 13):
            fins_de_mois.append((date(a, m + 1, 1) if m < 12 else date(a + 1, 1, 1)) - timedelta(days=1))
    releves = []
    for fin in fins_de_mois:
        for c in comptes_lus:
            etat_q, etat_e = _etat_compte_a(txs, c.code, fin, params.convention)
            for titre, q in sorted(etat_q.items()):
                if abs(q) > 1e-9:
                    releves.append([c.code, fin, titre, q])
            releves.append([c.code, fin, ENCAISSE, round(etat_e, 2)])
    ws = wb["Relevés"]
    for i, l in enumerate(releves, start=2):
        for j, v in enumerate(l, start=1):
            ws.cell(row=i, column=j, value=v)
    # un écart volontaire pour montrer la détection
    ws.cell(row=len(releves) + 2, column=1, value="RCM-1002")
    ws.cell(row=len(releves) + 2, column=2, value=date(2025, 12, 31))
    ws.cell(row=len(releves) + 2, column=3, value="VTI")
    ws.cell(row=len(releves) + 2, column=4, value=60)
    ws.cell(row=len(releves) + 2, column=5, value="Écart volontaire pour la démo (le relevé dit 60, le calcul dit 50)")

    ws = wb["1099"]
    ligne = 2
    for a in (2024, 2025):
        for c in comptes_lus:
            revs = [r for r in res.revenus if r.compte == c.code and r.date.year == a]
            disp = [d for d in res.dispositions if d.compte == c.code and d.date.year == a]
            vals = [c.code, a,
                    round(sum(r.brut_usd for r in revs if r.categorie == "Dividende"), 2),
                    round(sum(r.retenue_usd for r in revs), 2),
                    round(sum(r.brut_usd for r in revs if r.categorie == "Intérêt"), 2),
                    round(sum(d.produit_usd for d in disp), 2)]
            for j, v in enumerate(vals, start=1):
                ws.cell(row=ligne, column=j, value=v)
            ligne += 1
    wb.save(chemin)
    return chemin, chemin_taux


def _etat_compte_a(txs: list[Transaction], code: str, fin: date, convention: str):
    """Quantités et encaisse d'un compte à une date, à partir des opérations brutes (démo seulement)."""
    from collections import defaultdict
    q = defaultdict(float)
    e = 0.0
    for t in sorted(txs, key=lambda t: (t.date_effective(convention), t.ordre)):
        d = t.date_effective(convention)
        if d > fin:
            break
        brut = t.montant_brut()
        if t.compte == code:
            if t.type == "SOLDE_OUVERTURE":
                if t.titre == ENCAISSE:
                    e += brut
                else:
                    q[t.titre] += t.quantite
            elif t.type == "ACHAT":
                q[t.titre] += t.quantite
                e -= brut + t.commission_usd
            elif t.type == "VENTE":
                q[t.titre] -= t.quantite
                e += brut - t.commission_usd
            elif t.type in ("DIVIDENDE", "INTERET"):
                e += brut - t.retenue_usd
            elif t.type == "REMB_CAPITAL":
                e += brut
            elif t.type == "FRAIS":
                e -= brut
            elif t.type == "FRACTIONNEMENT":
                q[t.titre] += t.quantite
            elif t.type == "TRANSFERT_TITRES":
                q[t.titre] -= t.quantite
            elif t.type == "TRANSFERT_ENCAISSE":
                e -= brut
            elif t.type == "DEPOT":
                e += brut
            elif t.type == "RETRAIT":
                e -= brut
        if t.contrepartie == code:
            if t.type == "TRANSFERT_TITRES":
                q[t.titre] += t.quantite
            elif t.type == "TRANSFERT_ENCAISSE":
                e += brut
    return q, e


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Crée le gabarit de saisie")
    ap.add_argument("--demo", action="store_true", help="Gabarit rempli de données fictives + taux fictifs")
    ap.add_argument("--dossier", type=Path, default=Path("."), help="Dossier de sortie")
    ap.add_argument("--client", default="", help="Nom du client (gabarit vide)")
    ap.add_argument("--annee", type=int, default=date.today().year)
    ap.add_argument("--comptes", nargs="*", default=["COMPTE-1"], help="Codes des comptes (gabarit vide)")
    args = ap.parse_args(argv)
    args.dossier.mkdir(parents=True, exist_ok=True)
    if args.demo:
        g, t = demo(args.dossier)
        print(f"Créé : {g}\nCréé : {t} (taux FICTIFS)")
        return 0
    chemin = args.dossier / "Gabarit_Client_vide.xlsx"
    creer_classeur(chemin, args.client, args.annee, [Compte(c) for c in args.comptes])
    print(f"Créé : {chemin}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
