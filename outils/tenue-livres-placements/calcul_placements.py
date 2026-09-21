#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Tenue de livres de placements détenus dans des comptes américains
au profit d'un contribuable canadien.

Lit un classeur client (un onglet par compte), attache le taux de change
de la Banque du Canada à chaque opération, tient le prix de base rajusté
(PBR) en coût moyen pondéré en CAD, calcule les gains et pertes en capital,
les revenus (dividendes, intérêts, remboursements de capital), l'impôt
étranger retenu, les rapprochements aux relevés et aux 1099, et prépare
les données du T1135.

Confidentialité : la seule communication externe est la requête de taux
à la Banque du Canada, qui ne contient que des dates. Aucune donnée client
ne quitte le poste. Aucune intelligence artificielle n'est utilisée.

Usage :
    python calcul_placements.py Gabarit_Client.xlsx --annee 2025
    python calcul_placements.py Gabarit_Client.xlsx --annee 2025 --taux taux.csv   (hors ligne)
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from bisect import bisect_right
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Constantes
# ---------------------------------------------------------------------------

VALET_URL = "https://www.bankofcanada.ca/valet/observations/{serie}/json"
SERIE_QUOTIDIENNE = "FXUSDCAD"   # taux quotidien USD -> CAD (1 USD = x CAD)
SERIE_ANNUELLE = "FXAUSDCAD"     # moyenne annuelle publiée par la BdC

ENCAISSE = "ENCAISSE"

ACHAT = "ACHAT"
VENTE = "VENTE"
DIVIDENDE = "DIVIDENDE"
INTERET = "INTERET"
REMB_CAPITAL = "REMB_CAPITAL"
FRAIS = "FRAIS"
FRACTIONNEMENT = "FRACTIONNEMENT"
TRANSFERT_TITRES = "TRANSFERT_TITRES"
TRANSFERT_ENCAISSE = "TRANSFERT_ENCAISSE"
DEPOT = "DEPOT"
RETRAIT = "RETRAIT"
AJUST_PBR = "AJUST_PBR"
SOLDE_OUVERTURE = "SOLDE_OUVERTURE"

TYPES_VALIDES = [
    ACHAT, VENTE, DIVIDENDE, INTERET, REMB_CAPITAL, FRAIS, FRACTIONNEMENT,
    TRANSFERT_TITRES, TRANSFERT_ENCAISSE, DEPOT, RETRAIT, AJUST_PBR, SOLDE_OUVERTURE,
]

CONVENTION_TRANSACTION = "TRANSACTION"
CONVENTION_REGLEMENT = "REGLEMENT"
METHODE_QUOTIDIEN = "QUOTIDIEN"
METHODE_MOYENNE = "MOYENNE_ANNUELLE"

FENETRE_PERTE_APPARENTE = 30      # jours avant et après la disposition
SEUIL_RETENUE_ALERTE = 0.155      # retenue > 15 % : excédent non créditable (20(11))
ECART_TAUX_MAX_JOURS = 7          # au-delà, on considère le taux manquant
TOLERANCE_RAPPROCHEMENT = 0.01

COLONNES_TX = [
    "Date transaction", "Date règlement", "Type", "Titre", "Description",
    "Quantité", "Prix USD", "Montant brut USD", "Commission USD", "Retenue USD",
    "Compte contrepartie", "PBR CAD (ouverture)", "Note",
]
COLONNES_COMPTES = ["Code", "Courtier", "Numéro de compte", "Description"]
COLONNES_RELEVES = ["Compte", "Date de fin", "Titre", "Quantité ou solde USD"]
COLONNES_1099 = [
    "Compte", "Année", "Dividendes bruts USD", "Impôt étranger retenu USD",
    "Intérêts USD", "Produits de disposition USD",
]


# ---------------------------------------------------------------------------
# Taux de change
# ---------------------------------------------------------------------------

class TauxManquant(Exception):
    pass


class Taux:
    """Taux quotidiens USD/CAD et moyennes annuelles de la Banque du Canada."""

    def __init__(self, quotidiens: dict[date, float], annuels: dict[int, float] | None = None):
        self.quotidiens = dict(quotidiens)
        self.dates = sorted(self.quotidiens)
        self.annuels = dict(annuels or {})
        self.utilises: dict[date, tuple[date, float]] = {}

    def _precedent(self, d: date) -> tuple[date, float]:
        i = bisect_right(self.dates, d) - 1
        if i < 0:
            raise TauxManquant(f"Aucun taux publié le {d} ou avant")
        dd = self.dates[i]
        return dd, self.quotidiens[dd]

    def au(self, d: date) -> float:
        """Taux du jour, ou du jour ouvrable précédent le plus rapproché (fin de semaine, férié)."""
        dd, v = self._precedent(d)
        if (d - dd).days > ECART_TAUX_MAX_JOURS:
            raise TauxManquant(f"Aucun taux publié entre le {dd} et le {d}")
        self.utilises[d] = (dd, v)
        return v

    def dernier_disponible(self, d: date) -> tuple[date, float]:
        """Dernier taux publié au plus tard le jour demandé, sans limite d'écart (évaluation de fin d'année)."""
        dd, v = self._precedent(d)
        self.utilises[d] = (dd, v)
        return dd, v

    def moyenne(self, annee: int) -> float:
        if annee in self.annuels:
            return self.annuels[annee]
        vals = [v for dd, v in self.quotidiens.items() if dd.year == annee]
        if not vals:
            raise TauxManquant(f"Aucun taux pour l'année {annee}")
        return sum(vals) / len(vals)

    def couvre(self, debut: date, fin: date) -> bool:
        if not self.dates:
            return False
        return self.dates[0] <= debut and (fin - self.dates[-1]).days <= ECART_TAUX_MAX_JOURS

    def fusionner(self, autre: "Taux") -> None:
        self.quotidiens.update(autre.quotidiens)
        self.annuels.update(autre.annuels)
        self.dates = sorted(self.quotidiens)


def telecharger_taux_bdc(debut: date, fin: date, timeout: int = 30) -> Taux:
    """Interroge l'API Valet de la Banque du Canada. Seules des dates sont transmises."""
    import requests  # importé ici pour que le mode hors ligne n'en dépende pas

    params = {"start_date": debut.isoformat(), "end_date": fin.isoformat()}
    r = requests.get(VALET_URL.format(serie=SERIE_QUOTIDIENNE), params=params, timeout=timeout)
    r.raise_for_status()
    quotidiens: dict[date, float] = {}
    for obs in r.json().get("observations", []):
        v = obs.get(SERIE_QUOTIDIENNE, {}).get("v")
        if v not in (None, ""):
            quotidiens[date.fromisoformat(obs["d"])] = float(v)

    annuels: dict[int, float] = {}
    try:
        p2 = {"start_date": date(debut.year, 1, 1).isoformat(), "end_date": fin.isoformat()}
        r2 = requests.get(VALET_URL.format(serie=SERIE_ANNUELLE), params=p2, timeout=timeout)
        r2.raise_for_status()
        for obs in r2.json().get("observations", []):
            v = obs.get(SERIE_ANNUELLE, {}).get("v")
            if v not in (None, ""):
                annuels[date.fromisoformat(obs["d"]).year] = float(v)
    except Exception:
        pass  # la moyenne sera calculée à partir des taux quotidiens
    return Taux(quotidiens, annuels)


def charger_taux_csv(chemin: Path) -> Taux:
    """Fichier CSV : colonnes date,serie,valeur (serie = FXUSDCAD ou FXAUSDCAD)."""
    quotidiens: dict[date, float] = {}
    annuels: dict[int, float] = {}
    with open(chemin, newline="", encoding="utf-8") as f:
        for ligne in csv.DictReader(f):
            serie = (ligne.get("serie") or SERIE_QUOTIDIENNE).strip()
            d = date.fromisoformat(ligne["date"].strip())
            v = float(ligne["valeur"])
            if serie == SERIE_ANNUELLE:
                annuels[d.year] = v
            else:
                quotidiens[d] = v
    return Taux(quotidiens, annuels)


def sauvegarder_taux_csv(chemin: Path, taux: Taux) -> None:
    with open(chemin, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["date", "serie", "valeur"])
        for d in taux.dates:
            w.writerow([d.isoformat(), SERIE_QUOTIDIENNE, f"{taux.quotidiens[d]:.4f}"])
        for a in sorted(taux.annuels):
            w.writerow([f"{a}-12-31", SERIE_ANNUELLE, f"{taux.annuels[a]:.4f}"])


def obtenir_taux(debut: date, fin: date, fichier_taux: Path | None = None,
                 cache: Path | None = None) -> Taux:
    """Mode hors ligne si fichier_taux est fourni; sinon cache local puis Banque du Canada."""
    if fichier_taux is not None:
        return charger_taux_csv(fichier_taux)
    existant = charger_taux_csv(cache) if cache is not None and cache.exists() else Taux({})
    if existant.couvre(debut, fin):
        return existant
    telecharge = telecharger_taux_bdc(debut, fin)
    existant.fusionner(telecharge)
    if cache is not None:
        sauvegarder_taux_csv(cache, existant)
    return existant


# ---------------------------------------------------------------------------
# Modèle de données
# ---------------------------------------------------------------------------

@dataclass
class Parametres:
    client: str = ""
    annee: int = date.today().year
    convention: str = CONVENTION_TRANSACTION
    methode_revenus: str = METHODE_QUOTIDIEN


@dataclass
class Compte:
    code: str
    courtier: str = ""
    numero: str = ""
    description: str = ""


@dataclass
class Transaction:
    compte: str
    type: str
    titre: str
    date_tx: date
    date_regl: date | None = None
    quantite: float = 0.0
    prix_usd: float = 0.0
    brut_usd: float | None = None
    commission_usd: float = 0.0
    retenue_usd: float = 0.0
    contrepartie: str = ""
    pbr_cad_ouverture: float | None = None
    description: str = ""
    note: str = ""
    origine: str = ""
    ordre: int = 0

    def date_effective(self, convention: str) -> date:
        if convention == CONVENTION_REGLEMENT and self.date_regl is not None:
            return self.date_regl
        return self.date_tx

    def montant_brut(self) -> float:
        if self.brut_usd is not None:
            return self.brut_usd
        return self.quantite * self.prix_usd


@dataclass
class LigneReleve:
    compte: str
    date_fin: date
    titre: str
    valeur: float
    origine: str = ""


@dataclass
class Ligne1099:
    compte: str
    annee: int
    dividendes_usd: float = 0.0
    retenue_usd: float = 0.0
    interets_usd: float = 0.0
    produits_usd: float = 0.0
    origine: str = ""


@dataclass
class Pool:
    """PBR regroupé par titre pour le contribuable (tous comptes confondus)."""
    quantite: float = 0.0
    cout_cad: float = 0.0
    cout_usd: float = 0.0

    def moyen_cad(self) -> float:
        return self.cout_cad / self.quantite if self.quantite else 0.0

    def moyen_usd(self) -> float:
        return self.cout_usd / self.quantite if self.quantite else 0.0


@dataclass
class EtatCompte:
    quantites: dict[str, float] = field(default_factory=lambda: defaultdict(float))
    encaisse: float = 0.0


@dataclass
class Alerte:
    niveau: str        # ERREUR, AVERTISSEMENT, INFO
    origine: str
    message: str


@dataclass
class Disposition:
    date: date
    compte: str
    titre: str
    quantite: float
    prix_usd: float
    commission_usd: float
    produit_usd: float
    taux: float
    produit_cad: float
    pbr_usd: float
    pbr_cad: float
    gain_usd: float
    gain_cad_avant: float
    perte_refusee_cad: float
    gain_cad: float
    note: str = ""
    origine: str = ""


@dataclass
class Revenu:
    date: date
    compte: str
    titre: str
    categorie: str
    brut_usd: float
    retenue_usd: float
    taux: float
    brut_cad: float
    retenue_cad: float
    note: str = ""
    origine: str = ""


@dataclass
class LigneJournal:
    date: date
    compte: str
    type: str
    titre: str
    quantite: float
    prix_usd: float
    brut_usd: float
    commission_usd: float
    retenue_usd: float
    taux: float | None
    montant_cad: float | None
    pool_quantite: float
    pool_cout_cad: float
    pool_moyen_cad: float
    pool_cout_usd: float
    encaisse_compte: float
    origine: str
    note: str = ""


@dataclass
class LigneRapprochementReleve:
    compte: str
    date_fin: date
    titre: str
    releve: float
    calcule: float
    ecart: float
    statut: str


@dataclass
class LigneRapprochement1099:
    compte: str
    annee: int
    element: str
    f1099: float
    calcule: float
    ecart: float
    statut: str


@dataclass
class LigneT1135:
    element: str
    description: str
    cout_max_cad: float
    cout_fin_cad: float
    revenu_cad: float
    gain_cad: float


@dataclass
class Resultat:
    params: Parametres
    comptes: list[Compte] = field(default_factory=list)
    alertes: list[Alerte] = field(default_factory=list)
    dispositions: list[Disposition] = field(default_factory=list)
    revenus: list[Revenu] = field(default_factory=list)
    frais: list[Revenu] = field(default_factory=list)
    journal: list[LigneJournal] = field(default_factory=list)
    pools_fin: dict[str, Pool] = field(default_factory=dict)
    etats_fin: dict[str, EtatCompte] = field(default_factory=dict)
    taux_fin_annee: tuple[date, float] | None = None
    rapprochement_releves: list[LigneRapprochementReleve] = field(default_factory=list)
    rapprochement_1099: list[LigneRapprochement1099] = field(default_factory=list)
    t1135: list[LigneT1135] = field(default_factory=list)
    taux_moyen: float | None = None
    descriptions: dict[str, str] = field(default_factory=dict)
    taux: Taux | None = None

    def erreurs(self) -> list[Alerte]:
        return [a for a in self.alertes if a.niveau == "ERREUR"]


# ---------------------------------------------------------------------------
# Moteur de calcul
# ---------------------------------------------------------------------------

def calculer(transactions: Iterable[Transaction], comptes: Iterable[Compte], taux: Taux,
             params: Parametres, releves: Iterable[LigneReleve] = (),
             lignes_1099: Iterable[Ligne1099] = ()) -> Resultat:
    comptes = list(comptes)
    res = Resultat(params=params, comptes=comptes, taux=taux)
    codes = {c.code for c in comptes}
    conv = params.convention
    annee = params.annee
    debut_annee = date(annee, 1, 1)
    fin_annee = date(annee, 12, 31)

    def alerte(niveau: str, origine: str, message: str) -> None:
        res.alertes.append(Alerte(niveau, origine, message))

    # --- validation et tri --------------------------------------------------
    txs: list[Transaction] = []
    for i, t in enumerate(transactions):
        t.ordre = i
        t.titre = (t.titre or "").strip().upper()
        t.type = (t.type or "").strip().upper()
        if t.type not in TYPES_VALIDES:
            alerte("ERREUR", t.origine, f"Type d'opération inconnu : « {t.type} »")
            continue
        if t.compte not in codes:
            alerte("ERREUR", t.origine, f"Compte inconnu : « {t.compte} »")
            continue
        if t.type in (TRANSFERT_TITRES, TRANSFERT_ENCAISSE) and t.contrepartie not in codes:
            alerte("ERREUR", t.origine, f"Compte contrepartie inconnu : « {t.contrepartie} »")
            continue
        if t.type in (TRANSFERT_TITRES, TRANSFERT_ENCAISSE) and t.contrepartie == t.compte:
            alerte("ERREUR", t.origine, "Transfert vers le même compte")
            continue
        if not t.titre:
            alerte("ERREUR", t.origine, "Titre manquant (utiliser ENCAISSE pour les mouvements de trésorerie)")
            continue
        if t.type in (ACHAT, VENTE, FRACTIONNEMENT, TRANSFERT_TITRES) and t.titre == ENCAISSE:
            alerte("ERREUR", t.origine, f"{t.type} sur ENCAISSE n'a pas de sens")
            continue
        if t.type in (ACHAT, VENTE, TRANSFERT_TITRES) and t.quantite <= 0:
            alerte("ERREUR", t.origine, f"{t.type} : la quantité doit être positive")
            continue
        if t.description and t.titre not in res.descriptions:
            res.descriptions[t.titre] = t.description
        txs.append(t)
    txs.sort(key=lambda t: (t.date_effective(conv), t.ordre))

    # --- pré-passe pour la perte apparente (quantités détenues par le client) ---
    deltas: dict[str, list[tuple[date, float]]] = defaultdict(list)
    for t in txs:
        d = t.date_effective(conv)
        if t.titre == ENCAISSE:
            continue
        if t.type in (ACHAT, SOLDE_OUVERTURE, FRACTIONNEMENT):
            deltas[t.titre].append((d, t.quantite))
        elif t.type == VENTE:
            deltas[t.titre].append((d, -t.quantite))

    def detenu_fin_de_jour(titre: str, d: date) -> float:
        return sum(q for dd, q in deltas[titre] if dd <= d)

    def achats_fenetre(titre: str, d: date) -> float:
        lo = d - timedelta(days=FENETRE_PERTE_APPARENTE)
        hi = d + timedelta(days=FENETRE_PERTE_APPARENTE)
        return sum(t.quantite for t in txs
                   if t.type == ACHAT and t.titre == titre and lo <= t.date_effective(conv) <= hi)

    # --- état ---------------------------------------------------------------
    pools: dict[str, Pool] = defaultdict(Pool)
    etats: dict[str, EtatCompte] = {c.code: EtatCompte() for c in comptes}
    max_cout: dict[str, float] = defaultdict(float)
    max_encaisse: dict[str, float] = defaultdict(float)
    revenu_par_titre: dict[str, float] = defaultdict(float)
    gain_par_titre: dict[str, float] = defaultdict(float)
    interets_par_compte: dict[str, float] = defaultdict(float)
    dans_annee_atteint = False

    def taux_du_jour(d: date, origine: str) -> float | None:
        try:
            return taux.au(d)
        except TauxManquant as e:
            alerte("ERREUR", origine, str(e))
            return None

    def taux_revenu(d: date, origine: str) -> float | None:
        if params.methode_revenus == METHODE_MOYENNE:
            try:
                return taux.moyenne(d.year)
            except TauxManquant as e:
                alerte("ERREUR", origine, str(e))
                return None
        return taux_du_jour(d, origine)

    def maj_max(d: date) -> None:
        if not (debut_annee <= d <= fin_annee):
            return
        for titre, p in pools.items():
            max_cout[titre] = max(max_cout[titre], p.cout_cad)
        try:
            _, tx = taux.dernier_disponible(d)
        except TauxManquant:
            return
        for code, e in etats.items():
            max_encaisse[code] = max(max_encaisse[code], e.encaisse * tx)

    # --- points de contrôle (relevés, ouverture et fermeture d'année) ------------
    points: list[tuple[date, Callable[[], None]]] = []

    def ouvrir_annee() -> None:
        for titre, p in pools.items():
            max_cout[titre] = p.cout_cad
        try:
            _, tx = taux.dernier_disponible(debut_annee - timedelta(days=1))
            for code, e in etats.items():
                max_encaisse[code] = e.encaisse * tx
        except TauxManquant:
            pass

    def cloturer_annee() -> None:
        res.pools_fin = {k: Pool(p.quantite, p.cout_cad, p.cout_usd) for k, p in pools.items()}
        res.etats_fin = {k: EtatCompte(defaultdict(float, e.quantites), e.encaisse) for k, e in etats.items()}
        try:
            res.taux_fin_annee = taux.dernier_disponible(fin_annee)
            if (fin_annee - res.taux_fin_annee[0]).days > ECART_TAUX_MAX_JOURS:
                alerte("AVERTISSEMENT", "Fin d'année",
                       f"Taux de fin d'année : dernier taux disponible le {res.taux_fin_annee[0]}")
        except TauxManquant as e:
            alerte("AVERTISSEMENT", "Fin d'année", str(e))

    def verifier_releve(r: LigneReleve) -> Callable[[], None]:
        def _verifier() -> None:
            if r.compte not in etats:
                alerte("ERREUR", r.origine, f"Relevé : compte inconnu « {r.compte} »")
                return
            e = etats[r.compte]
            titre = r.titre.strip().upper()
            calcule = e.encaisse if titre == ENCAISSE else e.quantites.get(titre, 0.0)
            ecart = r.valeur - calcule
            statut = "OK" if abs(ecart) <= TOLERANCE_RAPPROCHEMENT else "ÉCART"
            res.rapprochement_releves.append(
                LigneRapprochementReleve(r.compte, r.date_fin, titre, r.valeur, calcule, ecart, statut))
            if statut == "ÉCART":
                alerte("AVERTISSEMENT", r.origine,
                       f"Relevé {r.compte} au {r.date_fin} : {titre} selon relevé {r.valeur:,.2f}, "
                       f"calculé {calcule:,.2f}, écart {ecart:,.2f}")
        return _verifier

    points.append((debut_annee - timedelta(days=1), ouvrir_annee))
    points.append((fin_annee, cloturer_annee))
    for r in releves:
        points.append((r.date_fin, verifier_releve(r)))
    points.sort(key=lambda p: p[0])

    # --- traitement d'une opération ---------------------------------------------
    def traiter(t: Transaction, d: date) -> None:
        nonlocal dans_annee_atteint
        e = etats[t.compte]
        p = pools[t.titre] if t.titre != ENCAISSE else None
        brut = t.montant_brut()
        encaisse_avant = e.encaisse
        tx: float | None = None
        montant_cad: float | None = None
        note = t.note or ""

        if t.type == SOLDE_OUVERTURE:
            if t.titre == ENCAISSE:
                e.encaisse += brut
            else:
                if t.pbr_cad_ouverture is None:
                    alerte("ERREUR", t.origine, "Solde d'ouverture : PBR CAD (ouverture) manquant")
                    return
                p.quantite += t.quantite
                p.cout_usd += brut
                p.cout_cad += t.pbr_cad_ouverture
                e.quantites[t.titre] += t.quantite
                montant_cad = t.pbr_cad_ouverture

        elif t.type == ACHAT:
            tx = taux_du_jour(d, t.origine)
            if tx is None:
                return
            cout_usd = brut + t.commission_usd
            montant_cad = cout_usd * tx
            p.quantite += t.quantite
            p.cout_usd += cout_usd
            p.cout_cad += montant_cad
            e.quantites[t.titre] += t.quantite
            e.encaisse -= cout_usd

        elif t.type == VENTE:
            tx = taux_du_jour(d, t.origine)
            if tx is None:
                return
            if p.quantite + 1e-9 < t.quantite:
                alerte("ERREUR", t.origine,
                       f"Vente de {t.quantite:g} {t.titre} alors que le client en détient {p.quantite:g}")
                return
            if e.quantites[t.titre] + 1e-9 < t.quantite:
                alerte("AVERTISSEMENT", t.origine,
                       f"Vente de {t.quantite:g} {t.titre} dans {t.compte} qui n'en détient que "
                       f"{e.quantites[t.titre]:g} (transfert manquant ?)")
            produit_usd = brut - t.commission_usd
            produit_cad = brut * tx - t.commission_usd * tx
            pbr_cad = p.moyen_cad() * t.quantite
            pbr_usd = p.moyen_usd() * t.quantite
            gain_cad_avant = produit_cad - pbr_cad
            gain_usd = produit_usd - pbr_usd
            p.quantite -= t.quantite
            p.cout_cad -= pbr_cad
            p.cout_usd -= pbr_usd
            if abs(p.quantite) < 1e-9:
                p.quantite = 0.0
                p.cout_cad = 0.0
                p.cout_usd = 0.0
            e.quantites[t.titre] -= t.quantite
            e.encaisse += produit_usd

            refusee = 0.0
            if gain_cad_avant < 0:
                achetes = achats_fenetre(t.titre, d)
                detenu = detenu_fin_de_jour(t.titre, d + timedelta(days=FENETRE_PERTE_APPARENTE))
                if achetes > 0 and detenu > 0:
                    fraction = min(t.quantite, achetes, detenu) / t.quantite
                    refusee = -gain_cad_avant * fraction
                    p.cout_cad += refusee
                    note = (note + " " if note else "") + (
                        f"Perte apparente : {refusee:,.2f} CAD refusé ({fraction:.0%}) et ajouté au PBR")
                    alerte("AVERTISSEMENT", t.origine,
                           f"Perte apparente sur {t.titre} : {refusee:,.2f} CAD refusé, ajouté au PBR")
            gain_cad = gain_cad_avant + refusee
            montant_cad = produit_cad
            res.dispositions.append(Disposition(
                d, t.compte, t.titre, t.quantite, t.prix_usd, t.commission_usd, produit_usd, tx,
                produit_cad, pbr_usd, pbr_cad, gain_usd, gain_cad_avant, refusee, gain_cad, note, t.origine))
            if debut_annee <= d <= fin_annee:
                gain_par_titre[t.titre] += gain_cad

        elif t.type in (DIVIDENDE, INTERET):
            tx = taux_revenu(d, t.origine)
            if tx is None:
                return
            brut_cad = brut * tx
            retenue_cad = t.retenue_usd * tx
            montant_cad = brut_cad
            e.encaisse += brut - t.retenue_usd
            categorie = "Dividende" if t.type == DIVIDENDE else "Intérêt"
            if t.type == DIVIDENDE and brut > 0:
                ratio = t.retenue_usd / brut
                if ratio > SEUIL_RETENUE_ALERTE:
                    alerte("AVERTISSEMENT", t.origine,
                           f"Retenue de {ratio:.1%} sur {t.titre} : excédent au-delà de 15 % non "
                           f"créditable (W-8BEN manquant ? déduction 20(11) à considérer)")
                elif ratio == 0:
                    alerte("INFO", t.origine, f"Dividende {t.titre} sans retenue étrangère")
            res.revenus.append(Revenu(d, t.compte, t.titre, categorie, brut, t.retenue_usd, tx,
                                      brut_cad, retenue_cad, note, t.origine))
            if debut_annee <= d <= fin_annee:
                if t.titre == ENCAISSE:
                    interets_par_compte[t.compte] += brut_cad
                else:
                    revenu_par_titre[t.titre] += brut_cad

        elif t.type == REMB_CAPITAL:
            tx = taux_du_jour(d, t.origine)
            if tx is None:
                return
            montant_cad = brut * tx
            e.encaisse += brut
            gain_repute = 0.0
            if p.cout_cad - montant_cad < -1e-9:
                gain_repute = montant_cad - p.cout_cad
                montant_pbr = p.cout_cad
                note = (note + " " if note else "") + (
                    f"Remboursement supérieur au PBR : gain en capital réputé de {gain_repute:,.2f} CAD")
                alerte("AVERTISSEMENT", t.origine, note)
            else:
                montant_pbr = montant_cad
            p.cout_cad -= montant_pbr
            p.cout_usd = max(p.cout_usd - brut, 0.0)
            res.revenus.append(Revenu(d, t.compte, t.titre, "Remboursement de capital", brut, 0.0, tx,
                                      montant_cad, 0.0, note, t.origine))
            if gain_repute > 0:
                res.dispositions.append(Disposition(
                    d, t.compte, t.titre, 0.0, 0.0, 0.0, 0.0, tx, 0.0, 0.0, 0.0, 0.0,
                    gain_repute, 0.0, gain_repute, "Gain réputé : remboursement de capital > PBR", t.origine))
                if debut_annee <= d <= fin_annee:
                    gain_par_titre[t.titre] += gain_repute

        elif t.type == FRAIS:
            tx = taux_du_jour(d, t.origine)
            if tx is None:
                return
            montant_cad = brut * tx
            e.encaisse -= brut
            res.frais.append(Revenu(d, t.compte, t.titre, "Frais", brut, 0.0, tx, montant_cad, 0.0,
                                    note, t.origine))

        elif t.type == FRACTIONNEMENT:
            if p.quantite + t.quantite < -1e-9:
                alerte("ERREUR", t.origine, f"Fractionnement : quantité résultante négative pour {t.titre}")
                return
            p.quantite += t.quantite
            e.quantites[t.titre] += t.quantite

        elif t.type == TRANSFERT_TITRES:
            if e.quantites[t.titre] + 1e-9 < t.quantite:
                alerte("AVERTISSEMENT", t.origine,
                       f"Transfert de {t.quantite:g} {t.titre} depuis {t.compte} qui n'en détient que "
                       f"{e.quantites[t.titre]:g}")
            e.quantites[t.titre] -= t.quantite
            etats[t.contrepartie].quantites[t.titre] += t.quantite
            note = (note + " " if note else "") + f"Vers {t.contrepartie} ; PBR inchangé (même contribuable)"

        elif t.type == TRANSFERT_ENCAISSE:
            e.encaisse -= brut
            etats[t.contrepartie].encaisse += brut
            note = (note + " " if note else "") + f"Vers {t.contrepartie}"

        elif t.type == DEPOT:
            e.encaisse += brut

        elif t.type == RETRAIT:
            e.encaisse -= brut

        elif t.type == AJUST_PBR:
            tx = taux_du_jour(d, t.origine)
            if tx is None:
                return
            montant_cad = brut * tx
            p.cout_cad += montant_cad
            p.cout_usd += brut

        if p is not None and p.quantite < -1e-9:
            alerte("ERREUR", t.origine, f"Quantité négative pour {t.titre} après cette opération")
        if e.encaisse < -0.005 and encaisse_avant >= -0.005:
            alerte("AVERTISSEMENT", t.origine,
                   f"Encaisse de {t.compte} négative ({e.encaisse:,.2f} USD) après cette opération : "
                   f"dépôt, vente, intérêt ou solde d'ouverture manquant ?")

        res.journal.append(LigneJournal(
            d, t.compte, t.type, t.titre, t.quantite, t.prix_usd, brut, t.commission_usd,
            t.retenue_usd, tx, montant_cad,
            p.quantite if p else 0.0, p.cout_cad if p else 0.0, p.moyen_cad() if p else 0.0,
            p.cout_usd if p else 0.0, e.encaisse, t.origine, note))
        maj_max(d)

    # --- boucle principale --------------------------------------------------
    idx = 0
    for t in txs:
        d = t.date_effective(conv)
        while idx < len(points) and points[idx][0] < d:
            points[idx][1]()
            idx += 1
        traiter(t, d)
    while idx < len(points):
        points[idx][1]()
        idx += 1

    # --- rapprochement 1099 ---------------------------------------------------
    for l in lignes_1099:
        if l.compte not in etats:
            alerte("ERREUR", l.origine, f"1099 : compte inconnu « {l.compte} »")
            continue
        revs = [r for r in res.revenus if r.compte == l.compte and r.date.year == l.annee]
        disp = [x for x in res.dispositions if x.compte == l.compte and x.date.year == l.annee]
        calc = {
            "Dividendes bruts USD": sum(r.brut_usd for r in revs if r.categorie == "Dividende"),
            "Impôt étranger retenu USD": sum(r.retenue_usd for r in revs),
            "Intérêts USD": sum(r.brut_usd for r in revs if r.categorie == "Intérêt"),
            "Produits de disposition USD": sum(x.produit_usd for x in disp),
        }
        decl = {
            "Dividendes bruts USD": l.dividendes_usd,
            "Impôt étranger retenu USD": l.retenue_usd,
            "Intérêts USD": l.interets_usd,
            "Produits de disposition USD": l.produits_usd,
        }
        for k in calc:
            ecart = decl[k] - calc[k]
            statut = "OK" if abs(ecart) <= TOLERANCE_RAPPROCHEMENT else "ÉCART"
            res.rapprochement_1099.append(LigneRapprochement1099(l.compte, l.annee, k, decl[k], calc[k], ecart, statut))
            if statut == "ÉCART":
                alerte("AVERTISSEMENT", l.origine,
                       f"1099 {l.compte} {l.annee} : {k} déclaré {decl[k]:,.2f}, calculé {calc[k]:,.2f}")

    # --- T1135 ---------------------------------------------------------------
    for titre in sorted(set(pools) | set(max_cout)):
        p = res.pools_fin.get(titre, Pool())
        if max_cout[titre] <= 0 and p.cout_cad <= 0 and revenu_par_titre[titre] == 0 and gain_par_titre[titre] == 0:
            continue
        res.t1135.append(LigneT1135(titre, res.descriptions.get(titre, ""), max_cout[titre], p.cout_cad,
                                    revenu_par_titre[titre], gain_par_titre[titre]))
    tx_fin = res.taux_fin_annee[1] if res.taux_fin_annee else None
    for c in comptes:
        e = res.etats_fin.get(c.code, EtatCompte())
        fin_cad = e.encaisse * tx_fin if tx_fin is not None else 0.0
        res.t1135.append(LigneT1135(f"{ENCAISSE} {c.code}", f"Encaisse USD chez {c.courtier}".strip(),
                                    max_encaisse[c.code], fin_cad, interets_par_compte[c.code], 0.0))

    if params.methode_revenus == METHODE_MOYENNE:
        try:
            res.taux_moyen = taux.moyenne(annee)
        except TauxManquant:
            pass
    return res


# ---------------------------------------------------------------------------
# Lecture du classeur client
# ---------------------------------------------------------------------------

def _texte(v) -> str:
    return "" if v is None else str(v).strip()


def _nombre(v, defaut: float | None = 0.0) -> float | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return defaut
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace(" ", "").replace("$", "")
    if "," in s and "." not in s:
        s = s.replace(",", ".")
    else:
        s = s.replace(",", "")
    return float(s)


def _date(v) -> date | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Date illisible : « {s} »")


def _entetes(ws) -> dict[str, int]:
    return {_texte(c.value).lower(): i for i, c in enumerate(ws[1]) if _texte(c.value)}


def _col(entetes: dict[str, int], nom: str) -> int | None:
    return entetes.get(nom.lower())


def _lire_table(ws, colonnes: list[str]) -> list[tuple[int, dict[str, object]]]:
    ent = _entetes(ws)
    manquantes = [c for c in colonnes if _col(ent, c) is None]
    if manquantes:
        raise ValueError(f"Onglet « {ws.title} » : colonnes manquantes {manquantes}")
    lignes = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None or _texte(v) == "" for v in row):
            continue
        lignes.append((i, {c: row[_col(ent, c)] if _col(ent, c) < len(row) else None for c in colonnes}))
    return lignes


def lire_classeur(chemin: Path, annee: int | None = None):
    wb = load_workbook(chemin, data_only=True)
    params = Parametres()
    if "Paramètres" in wb.sheetnames:
        for row in wb["Paramètres"].iter_rows(min_row=1, values_only=True):
            if not row or row[0] is None:
                continue
            cle = _texte(row[0]).lower()
            val = row[1] if len(row) > 1 else None
            if cle == "client":
                params.client = _texte(val)
            elif cle in ("année", "annee"):
                params.annee = int(_nombre(val, float(date.today().year)))
            elif cle.startswith("convention"):
                v = _texte(val).upper()
                params.convention = CONVENTION_REGLEMENT if v.startswith("R") else CONVENTION_TRANSACTION
            elif cle.startswith("taux pour les revenus") or cle.startswith("méthode"):
                v = _texte(val).upper()
                params.methode_revenus = METHODE_MOYENNE if v.startswith("MOY") else METHODE_QUOTIDIEN
    if annee is not None:
        params.annee = annee

    comptes: list[Compte] = []
    for i, r in _lire_table(wb["Comptes"], COLONNES_COMPTES):
        code = _texte(r["Code"])
        if code:
            comptes.append(Compte(code, _texte(r["Courtier"]), _texte(r["Numéro de compte"]), _texte(r["Description"])))

    transactions: list[Transaction] = []
    erreurs: list[Alerte] = []
    for c in comptes:
        if c.code not in wb.sheetnames:
            erreurs.append(Alerte("ERREUR", "Comptes", f"Aucun onglet nommé « {c.code} » pour ce compte"))
            continue
        for i, r in _lire_table(wb[c.code], COLONNES_TX):
            origine = f"{c.code}!L{i}"
            try:
                d_tx = _date(r["Date transaction"])
                if d_tx is None:
                    raise ValueError("Date transaction manquante")
                transactions.append(Transaction(
                    compte=c.code,
                    type=_texte(r["Type"]).upper(),
                    titre=_texte(r["Titre"]).upper(),
                    date_tx=d_tx,
                    date_regl=_date(r["Date règlement"]),
                    quantite=_nombre(r["Quantité"]) or 0.0,
                    prix_usd=_nombre(r["Prix USD"]) or 0.0,
                    brut_usd=_nombre(r["Montant brut USD"], None),
                    commission_usd=_nombre(r["Commission USD"]) or 0.0,
                    retenue_usd=_nombre(r["Retenue USD"]) or 0.0,
                    contrepartie=_texte(r["Compte contrepartie"]),
                    pbr_cad_ouverture=_nombre(r["PBR CAD (ouverture)"], None),
                    description=_texte(r["Description"]),
                    note=_texte(r["Note"]),
                    origine=origine,
                ))
            except Exception as e:  # ligne illisible : on la signale sans arrêter le traitement
                erreurs.append(Alerte("ERREUR", origine, f"Ligne illisible : {e}"))

    releves: list[LigneReleve] = []
    if "Relevés" in wb.sheetnames:
        for i, r in _lire_table(wb["Relevés"], COLONNES_RELEVES):
            try:
                releves.append(LigneReleve(_texte(r["Compte"]), _date(r["Date de fin"]), _texte(r["Titre"]).upper(),
                                           _nombre(r["Quantité ou solde USD"]) or 0.0, f"Relevés!L{i}"))
            except Exception as e:
                erreurs.append(Alerte("ERREUR", f"Relevés!L{i}", f"Ligne illisible : {e}"))

    lignes_1099: list[Ligne1099] = []
    if "1099" in wb.sheetnames:
        for i, r in _lire_table(wb["1099"], COLONNES_1099):
            try:
                lignes_1099.append(Ligne1099(
                    _texte(r["Compte"]), int(_nombre(r["Année"])),
                    _nombre(r["Dividendes bruts USD"]) or 0.0, _nombre(r["Impôt étranger retenu USD"]) or 0.0,
                    _nombre(r["Intérêts USD"]) or 0.0, _nombre(r["Produits de disposition USD"]) or 0.0,
                    f"1099!L{i}"))
            except Exception as e:
                erreurs.append(Alerte("ERREUR", f"1099!L{i}", f"Ligne illisible : {e}"))

    return params, comptes, transactions, releves, lignes_1099, erreurs


# ---------------------------------------------------------------------------
# Écriture du classeur de sortie
# ---------------------------------------------------------------------------

_ENTETE_FILL = PatternFill("solid", fgColor="1F3864")
_ENTETE_FONT = Font(bold=True, color="FFFFFF")
_ECART_FILL = PatternFill("solid", fgColor="F8CBAD")
_OK_FILL = PatternFill("solid", fgColor="C6E0B4")
FMT_MONTANT = "#,##0.00;[Red]-#,##0.00"
FMT_QTE = "#,##0.####"
FMT_TAUX = "0.0000"
FMT_DATE = "yyyy-mm-dd"


def _ecrire_table(ws, entetes: list[str], lignes: list[list], formats: dict[int, str] | None = None,
                  ligne_depart: int = 1, largeurs: dict[int, float] | None = None) -> None:
    formats = formats or {}
    for j, h in enumerate(entetes, start=1):
        c = ws.cell(row=ligne_depart, column=j, value=h)
        c.fill = _ENTETE_FILL
        c.font = _ENTETE_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, ligne in enumerate(lignes, start=ligne_depart + 1):
        for j, v in enumerate(ligne, start=1):
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, (date, datetime)):
                c.number_format = FMT_DATE
            elif j in formats and isinstance(v, (int, float)):
                c.number_format = formats[j]
    ws.freeze_panes = ws.cell(row=ligne_depart + 1, column=1)
    for j in range(1, len(entetes) + 1):
        largeur = (largeurs or {}).get(j, max(12, min(45, len(entetes[j - 1]) + 4)))
        ws.column_dimensions[get_column_letter(j)].width = largeur


def _r(v: float | None) -> float | None:
    return None if v is None else round(v, 2)


def ecrire_sortie(res: Resultat, chemin: Path) -> None:
    p = res.params
    annee = p.annee
    dans_annee = lambda d: d.year == annee  # noqa: E731

    wb = Workbook()

    # --- Sommaire -------------------------------------------------------------
    ws = wb.active
    ws.title = "Sommaire"
    disp = [x for x in res.dispositions if dans_annee(x.date)]
    revs = [x for x in res.revenus if dans_annee(x.date)]
    frais = [x for x in res.frais if dans_annee(x.date)]
    div = [x for x in revs if x.categorie == "Dividende"]
    inte = [x for x in revs if x.categorie == "Intérêt"]
    rc = [x for x in revs if x.categorie == "Remboursement de capital"]
    gains = [x.gain_cad for x in disp if x.gain_cad > 0]
    pertes = [x.gain_cad for x in disp if x.gain_cad < 0]

    ws["A1"] = f"Placements américains : sommaire fiscal {annee}"
    ws["A1"].font = Font(bold=True, size=14)
    infos = [
        ("Client", p.client),
        ("Année d'imposition", annee),
        ("Convention de date", "Date de règlement" if p.convention == CONVENTION_REGLEMENT else "Date de transaction"),
        ("Taux pour les revenus", "Moyenne annuelle BdC" if p.methode_revenus == METHODE_MOYENNE else "Taux quotidien BdC"),
        ("Taux moyen annuel utilisé", res.taux_moyen),
        ("Taux de fin d'année (date, taux)", f"{res.taux_fin_annee[0]} : {res.taux_fin_annee[1]:.4f}" if res.taux_fin_annee else ""),
        ("Généré le", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Erreurs", len(res.erreurs())),
        ("Avertissements", len([a for a in res.alertes if a.niveau == "AVERTISSEMENT"])),
    ]
    for i, (k, v) in enumerate(infos, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=v)
        if k.startswith("Taux moyen") and isinstance(v, float):
            c.number_format = FMT_TAUX
    debut = len(infos) + 4
    lignes = [
        ["Produits de disposition (nets de commissions)", _r(sum(x.produit_cad for x in disp)), _r(sum(x.produit_usd for x in disp))],
        ["PBR des titres disposés", _r(sum(x.pbr_cad for x in disp)), _r(sum(x.pbr_usd for x in disp))],
        ["Gains en capital (avant pertes)", _r(sum(gains)), None],
        ["Pertes en capital (après pertes apparentes)", _r(sum(pertes)), None],
        ["Pertes apparentes refusées (ajoutées au PBR)", _r(sum(x.perte_refusee_cad for x in disp)), None],
        ["Gain (perte) en capital net : annexe 3 / annexe G", _r(sum(x.gain_cad for x in disp)), _r(sum(x.gain_usd for x in disp))],
        ["Dividendes bruts (revenu étranger)", _r(sum(x.brut_cad for x in div)), _r(sum(x.brut_usd for x in div))],
        ["Intérêts bruts (revenu étranger)", _r(sum(x.brut_cad for x in inte)), _r(sum(x.brut_usd for x in inte))],
        ["Remboursements de capital (réduisent le PBR)", _r(sum(x.brut_cad for x in rc)), _r(sum(x.brut_usd for x in rc))],
        ["Impôt étranger retenu : T2209 / TP-772", _r(sum(x.retenue_cad for x in revs)), _r(sum(x.retenue_usd for x in revs))],
        ["Frais (déductibilité à évaluer, 20(1)(bb))", _r(sum(x.brut_cad for x in frais)), _r(sum(x.brut_usd for x in frais))],
    ]
    _ecrire_table(ws, ["Élément", "CAD", "USD"], lignes, {2: FMT_MONTANT, 3: FMT_MONTANT}, ligne_depart=debut,
                  largeurs={1: 55, 2: 18, 3: 18})
    ws.freeze_panes = None
    ws.column_dimensions["B"].width = 26

    # --- Alertes ----------------------------------------------------------------
    ws = wb.create_sheet("Alertes")
    ordre = {"ERREUR": 0, "AVERTISSEMENT": 1, "INFO": 2}
    lignes = [[a.niveau, a.origine, a.message] for a in sorted(res.alertes, key=lambda a: ordre.get(a.niveau, 9))]
    _ecrire_table(ws, ["Niveau", "Origine (onglet!ligne)", "Message"], lignes, largeurs={1: 16, 2: 22, 3: 110})
    for i, a in enumerate(sorted(res.alertes, key=lambda a: ordre.get(a.niveau, 9)), start=2):
        if a.niveau == "ERREUR":
            ws.cell(row=i, column=1).fill = _ECART_FILL

    # --- Dispositions -------------------------------------------------------------
    ws = wb.create_sheet("Dispositions")
    lignes = [[x.date, x.compte, x.titre, x.quantite, _r(x.prix_usd), _r(x.commission_usd), _r(x.produit_usd),
               round(x.taux, 4), _r(x.produit_cad), _r(x.pbr_usd), _r(x.pbr_cad), _r(x.gain_usd),
               _r(x.gain_cad_avant), _r(x.perte_refusee_cad), _r(x.gain_cad), x.note, x.origine] for x in disp]
    _ecrire_table(ws, ["Date", "Compte", "Titre", "Quantité", "Prix USD", "Commission USD", "Produit net USD",
                       "Taux BdC", "Produit net CAD", "PBR USD", "PBR CAD", "Gain (perte) USD (indicatif)",
                       "Gain (perte) CAD avant perte apparente", "Perte apparente refusée CAD",
                       "Gain (perte) CAD imposable", "Note", "Origine"],
                  lignes, {4: FMT_QTE, 5: FMT_MONTANT, 6: FMT_MONTANT, 7: FMT_MONTANT, 8: FMT_TAUX, 9: FMT_MONTANT,
                           10: FMT_MONTANT, 11: FMT_MONTANT, 12: FMT_MONTANT, 13: FMT_MONTANT, 14: FMT_MONTANT,
                           15: FMT_MONTANT}, largeurs={16: 60})

    # --- Revenus --------------------------------------------------------------------
    ws = wb.create_sheet("Revenus")
    lignes = [[x.date, x.compte, x.titre, x.categorie, _r(x.brut_usd), _r(x.retenue_usd), _r(x.brut_usd - x.retenue_usd),
               round(x.taux, 4), _r(x.brut_cad), _r(x.retenue_cad), x.note, x.origine] for x in revs]
    _ecrire_table(ws, ["Date", "Compte", "Titre", "Catégorie", "Brut USD", "Retenue USD", "Net USD", "Taux",
                       "Brut CAD", "Retenue CAD", "Note", "Origine"],
                  lignes, {5: FMT_MONTANT, 6: FMT_MONTANT, 7: FMT_MONTANT, 8: FMT_TAUX, 9: FMT_MONTANT, 10: FMT_MONTANT},
                  largeurs={11: 50})
    # sous-totaux par titre
    ligne = len(lignes) + 3
    ws.cell(row=ligne, column=1, value="Sous-totaux par titre et catégorie").font = Font(bold=True)
    agg: dict[tuple[str, str], list[float]] = defaultdict(lambda: [0.0, 0.0, 0.0, 0.0])
    for x in revs:
        a = agg[(x.titre, x.categorie)]
        a[0] += x.brut_usd
        a[1] += x.retenue_usd
        a[2] += x.brut_cad
        a[3] += x.retenue_cad
    st = [[t, c, _r(v[0]), _r(v[1]), _r(v[2]), _r(v[3])] for (t, c), v in sorted(agg.items())]
    _ecrire_table(ws, ["Titre", "Catégorie", "Brut USD", "Retenue USD", "Brut CAD", "Retenue CAD"], st,
                  {3: FMT_MONTANT, 4: FMT_MONTANT, 5: FMT_MONTANT, 6: FMT_MONTANT}, ligne_depart=ligne + 1)
    ws.freeze_panes = "A2"

    # --- Par titre -----------------------------------------------------------------------
    ws = wb.create_sheet("Par titre", 2)
    ws["A1"] = f"Revenus, retenues et gains par placement pour {annee} (CAD et USD)"
    ws["A1"].font = Font(bold=True)
    par_titre: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for x in revs:
        cle = {"Dividende": "div", "Intérêt": "int", "Remboursement de capital": "rc"}[x.categorie]
        par_titre[x.titre][cle + "_usd"] += x.brut_usd
        par_titre[x.titre][cle + "_cad"] += x.brut_cad
        par_titre[x.titre]["ret_usd"] += x.retenue_usd
        par_titre[x.titre]["ret_cad"] += x.retenue_cad
    for x in disp:
        par_titre[x.titre]["gain_usd"] += x.gain_usd
        par_titre[x.titre]["gain_cad"] += x.gain_cad
        par_titre[x.titre]["refusee_cad"] += x.perte_refusee_cad
        par_titre[x.titre]["nb_disp"] += 1
    lignes = []
    for titre in sorted(par_titre):
        v = par_titre[titre]
        lignes.append([titre, res.descriptions.get(titre, ""),
                       _r(v["div_cad"]), _r(v["div_usd"]), _r(v["int_cad"]), _r(v["int_usd"]),
                       _r(v["rc_cad"]), _r(v["rc_usd"]), _r(v["ret_cad"]), _r(v["ret_usd"]),
                       int(v["nb_disp"]), _r(v["gain_cad"]), _r(v["gain_usd"]), _r(v["refusee_cad"])])
    if lignes:
        total = ["Total", ""] + [_r(sum(l[j] or 0 for l in lignes)) for j in range(2, 14)]
        total[10] = int(total[10])
        lignes.append(total)
    _ecrire_table(ws, ["Titre", "Description", "Dividendes CAD", "Dividendes USD", "Intérêts CAD", "Intérêts USD",
                       "Remb. de capital CAD", "Remb. de capital USD", "Impôt étranger retenu CAD",
                       "Impôt étranger retenu USD", "Nb de dispositions", "Gain (perte) en capital CAD",
                       "Gain (perte) USD (indicatif)", "Perte apparente refusée CAD"],
                  lignes, {j: FMT_MONTANT for j in range(3, 15) if j != 11}, ligne_depart=3, largeurs={2: 35})
    if lignes:
        for j in range(1, 15):
            ws.cell(row=len(lignes) + 3, column=j).font = Font(bold=True)
    ws.freeze_panes = "C4"

    # --- Frais -----------------------------------------------------------------------
    ws = wb.create_sheet("Frais")
    lignes = [[x.date, x.compte, _r(x.brut_usd), round(x.taux, 4), _r(x.brut_cad), x.note, x.origine] for x in frais]
    _ecrire_table(ws, ["Date", "Compte", "USD", "Taux", "CAD", "Note", "Origine"], lignes,
                  {3: FMT_MONTANT, 4: FMT_TAUX, 5: FMT_MONTANT}, largeurs={6: 50})

    # --- Positions ---------------------------------------------------------------------
    ws = wb.create_sheet("Positions")
    ws["A1"] = f"PBR regroupé par titre au 31 décembre {annee} (tous comptes, même contribuable)"
    ws["A1"].font = Font(bold=True)
    lignes = [[t, res.descriptions.get(t, ""), pl.quantite, _r(pl.cout_usd), _r(pl.moyen_usd()), _r(pl.cout_cad), _r(pl.moyen_cad())]
              for t, pl in sorted(res.pools_fin.items()) if abs(pl.quantite) > 1e-9 or abs(pl.cout_cad) > 0.005]
    _ecrire_table(ws, ["Titre", "Description", "Quantité", "PBR total USD", "PBR moyen USD", "PBR total CAD", "PBR moyen CAD"],
                  lignes, {3: FMT_QTE, 4: FMT_MONTANT, 5: FMT_MONTANT, 6: FMT_MONTANT, 7: FMT_MONTANT}, ligne_depart=3,
                  largeurs={2: 35})
    ligne = len(lignes) + 6
    ws.cell(row=ligne, column=1, value=f"Détail par compte au 31 décembre {annee}").font = Font(bold=True)
    det = []
    for code, e in sorted(res.etats_fin.items()):
        for t, q in sorted(e.quantites.items()):
            if abs(q) > 1e-9:
                pl = res.pools_fin.get(t, Pool())
                det.append([code, t, q, _r(pl.moyen_cad() * q), _r(pl.moyen_usd() * q)])
        det.append([code, ENCAISSE, _r(e.encaisse), _r(e.encaisse * res.taux_fin_annee[1]) if res.taux_fin_annee else None, _r(e.encaisse)])
    _ecrire_table(ws, ["Compte", "Titre", "Quantité / solde USD", "PBR ou valeur CAD (part du compte)", "PBR ou solde USD"],
                  det, {3: FMT_QTE, 4: FMT_MONTANT, 5: FMT_MONTANT}, ligne_depart=ligne + 1)
    ws.freeze_panes = None

    # --- Journal -----------------------------------------------------------------------
    ws = wb.create_sheet("Journal")
    lignes = [[j.date, j.compte, j.type, j.titre, j.quantite or None, _r(j.prix_usd) or None, _r(j.brut_usd),
               _r(j.commission_usd) or None, _r(j.retenue_usd) or None, round(j.taux, 4) if j.taux is not None else None,
               _r(j.montant_cad), j.pool_quantite, _r(j.pool_cout_cad), _r(j.pool_moyen_cad), _r(j.pool_cout_usd),
               _r(j.encaisse_compte), j.note, j.origine] for j in res.journal]
    _ecrire_table(ws, ["Date effective", "Compte", "Type", "Titre", "Quantité", "Prix USD", "Montant brut USD",
                       "Commission USD", "Retenue USD", "Taux BdC", "Montant CAD", "Pool : quantité après",
                       "Pool : PBR CAD après", "Pool : PBR moyen CAD", "Pool : PBR USD après", "Encaisse du compte USD après",
                       "Note", "Origine"],
                  lignes, {5: FMT_QTE, 6: FMT_MONTANT, 7: FMT_MONTANT, 8: FMT_MONTANT, 9: FMT_MONTANT, 10: FMT_TAUX,
                           11: FMT_MONTANT, 12: FMT_QTE, 13: FMT_MONTANT, 14: FMT_MONTANT, 15: FMT_MONTANT, 16: FMT_MONTANT},
                  largeurs={17: 60})

    # --- Rapprochements ----------------------------------------------------------------
    ws = wb.create_sheet("Rapprochement relevés")
    lignes = [[x.compte, x.date_fin, x.titre, x.releve, x.calcule, _r(x.ecart), x.statut]
              for x in sorted(res.rapprochement_releves, key=lambda x: (x.compte, x.date_fin, x.titre))]
    _ecrire_table(ws, ["Compte", "Date de fin", "Titre", "Selon relevé", "Calculé", "Écart", "Statut"], lignes,
                  {4: FMT_QTE, 5: FMT_QTE, 6: FMT_QTE})
    for i, l in enumerate(lignes, start=2):
        ws.cell(row=i, column=7).fill = _OK_FILL if l[6] == "OK" else _ECART_FILL

    ws = wb.create_sheet("Rapprochement 1099")
    lignes = [[x.compte, x.annee, x.element, _r(x.f1099), _r(x.calcule), _r(x.ecart), x.statut] for x in res.rapprochement_1099]
    _ecrire_table(ws, ["Compte", "Année", "Élément", "Selon 1099", "Calculé (USD)", "Écart", "Statut"], lignes,
                  {4: FMT_MONTANT, 5: FMT_MONTANT, 6: FMT_MONTANT}, largeurs={3: 32})
    for i, l in enumerate(lignes, start=2):
        ws.cell(row=i, column=7).fill = _OK_FILL if l[6] == "OK" else _ECART_FILL

    # --- T1135 ------------------------------------------------------------------------
    ws = wb.create_sheet("T1135")
    ws["A1"] = ("Données pour le T1135 (méthode détaillée). La catégorie (1 : fonds, 2 : actions de sociétés "
                "non résidentes, 4 : fiducies non résidentes, etc.) reste à déterminer par titre. "
                "Le seuil de 100 000 $ CAD s'évalue sur le coût total.")
    lignes = [[x.element, x.description, _r(x.cout_max_cad), _r(x.cout_fin_cad), _r(x.revenu_cad), _r(x.gain_cad)] for x in res.t1135]
    _ecrire_table(ws, ["Bien", "Description", "Coût maximal dans l'année CAD", "Coût à la fin de l'année CAD",
                       "Revenu brut CAD", "Gain (perte) CAD"], lignes,
                  {3: FMT_MONTANT, 4: FMT_MONTANT, 5: FMT_MONTANT, 6: FMT_MONTANT}, ligne_depart=3, largeurs={2: 35})
    tot = len(lignes) + 4
    ws.cell(row=tot, column=1, value="Total").font = Font(bold=True)
    for j in (3, 4, 5, 6):
        c = ws.cell(row=tot, column=j, value=_r(sum(l[j - 1] or 0 for l in lignes)))
        c.number_format = FMT_MONTANT
        c.font = Font(bold=True)
    ws.freeze_panes = None

    # --- Taux utilisés -----------------------------------------------------------------
    ws = wb.create_sheet("Taux")
    lignes = []
    if res.taux is not None:
        for d, (dd, v) in sorted(res.taux.utilises.items()):
            lignes.append([d, dd, round(v, 4), "" if d == dd else "Jour ouvrable précédent"])
    _ecrire_table(ws, ["Date de l'opération", "Date du taux BdC", "Taux USD/CAD", "Remarque"], lignes, {3: FMT_TAUX},
                  largeurs={4: 28})

    wb.save(chemin)


# ---------------------------------------------------------------------------
# Ligne de commande
# ---------------------------------------------------------------------------

def executer(classeur: Path, annee: int | None, sortie: Path | None, fichier_taux: Path | None,
             cache: Path | None) -> Resultat:
    params, comptes, transactions, releves, lignes_1099, erreurs_lecture = lire_classeur(classeur, annee)
    if not transactions:
        raise SystemExit("Aucune opération trouvée dans le classeur.")
    dates = [t.date_effective(params.convention) for t in transactions]
    debut = min(dates) - timedelta(days=45)
    fin = max(max(dates), date(params.annee, 12, 31))
    fin = min(fin, date.today())
    taux = obtenir_taux(debut, fin, fichier_taux, cache)
    res = calculer(transactions, comptes, taux, params, releves, lignes_1099)
    res.alertes = erreurs_lecture + res.alertes
    if sortie is None:
        nom_client = re.sub(r"[^A-Za-z0-9À-ÿ]+", "_", params.client).strip("_") or "Client"
        sortie = classeur.with_name(f"Sortie_{nom_client}_{params.annee}.xlsx")
    ecrire_sortie(res, sortie)
    return res


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Tenue de livres de placements américains (contribuable canadien)")
    ap.add_argument("classeur", type=Path, help="Classeur client (gabarit rempli)")
    ap.add_argument("--annee", type=int, default=None, help="Année d'imposition (sinon l'onglet Paramètres)")
    ap.add_argument("--sortie", type=Path, default=None, help="Fichier Excel de sortie")
    ap.add_argument("--taux", type=Path, default=None, help="Fichier CSV de taux (mode hors ligne, aucune connexion)")
    ap.add_argument("--cache", type=Path, default=None,
                    help="Cache local des taux BdC (défaut : taux_bdc_cache.csv à côté du classeur)")
    args = ap.parse_args(argv)
    cache = args.cache if args.cache is not None else args.classeur.with_name("taux_bdc_cache.csv")
    res = executer(args.classeur, args.annee, args.sortie, args.taux, cache)
    n_err = len(res.erreurs())
    n_av = len([a for a in res.alertes if a.niveau == "AVERTISSEMENT"])
    print(f"Terminé : {len(res.journal)} opérations traitées, {n_err} erreur(s), {n_av} avertissement(s).")
    for a in res.alertes:
        if a.niveau != "INFO":
            print(f"  [{a.niveau}] {a.origine} : {a.message}")
    return 1 if n_err else 0


if __name__ == "__main__":
    sys.exit(main())
